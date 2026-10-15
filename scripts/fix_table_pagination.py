# -*- coding: utf-8 -*-
"""docx 分页修复（article-format 生成后的固定后处理）。

四项修复：
1. 长表格分页——所有表格行加 w:cantSplit（行不跨页断开），表头行加
   w:tblHeader（跨页重复表头）。不修的话，表格行会在页边界被切成两半、
   续页顶部出现半截行且无表头。
2. 加粗小标题与标签行防孤行——整段加粗且较短的段落（如「LOOK 01 · …」标题行）、
   以及以全角冒号结尾的短标签行（如「英文版：」）加 w:keepNext（与下段同页）。
   不修的话，标题/标签会孤零零留在页尾、内容被推到下一页。
3. 表头行与首数据行同页——表头行单元格段落加 w:keepNext。不修的话，表头行
   （连同其前面的章节标题）可能孤立在页尾、数据行全部落到下一页，出现
   「页尾只有一行表头」的空视觉。
4. 「镜 N 操作顺序」短段防切半——仅段首加粗的操作顺序段加 w:keepLines
   （段中不分页）。不修的话，该段可能在页边界被从中间切断、标题块贴页尾
   （脚本类长文档实测）。

用法：python fix_table_pagination.py <docx路径>
     python fix_table_pagination.py --selftest
md2docx/article_style 重跑生成新 docx 后，对本脚本重跑一次即可。
"""
import sys

from docx import Document
from docx.oxml import OxmlElement

W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
TRPR_AFTER_CANTSPLIT = (
    "w:trHeight",
    "w:tblHeader",
    "w:tblCellSpacing",
    "w:jc",
    "w:hidden",
    "w:ins",
    "w:del",
    "w:trPrChange",
)
TRPR_AFTER_TBLHEADER = (
    "w:tblCellSpacing",
    "w:jc",
    "w:hidden",
    "w:ins",
    "w:del",
    "w:trPrChange",
)

BOLD_TITLE_MAX_LEN = 200
LABEL_MAX_LEN = 50
OP_LINE_MAX_LEN = 300


def fix_tables(doc) -> int:
    fixed_rows = 0
    for table in doc.tables:
        for i, row in enumerate(table.rows):
            tr = row._tr
            trPr = tr.get_or_add_trPr()
            if trPr.find(f"{W_NS}cantSplit") is None:
                trPr.insert_element_before(OxmlElement("w:cantSplit"), *TRPR_AFTER_CANTSPLIT)
                fixed_rows += 1
            if i == 0 and trPr.find(f"{W_NS}tblHeader") is None:
                trPr.insert_element_before(OxmlElement("w:tblHeader"), *TRPR_AFTER_TBLHEADER)
    return fixed_rows


def fix_heading_orphans(doc) -> int:
    fixed = 0
    for p in doc.paragraphs:
        text = p.text.strip()
        runs = [r for r in p.runs if r.text.strip()]
        if not text or p.paragraph_format.keep_with_next:
            continue
        is_bold_title = runs and len(text) <= BOLD_TITLE_MAX_LEN and all(r.bold for r in runs)
        is_label = len(text) <= LABEL_MAX_LEN and text.endswith("：")
        if is_bold_title or is_label:
            p.paragraph_format.keep_with_next = True
            fixed += 1
    return fixed


def fix_header_keepnext(doc) -> int:
    fixed = 0
    for table in doc.tables:
        if len(table.rows) < 2:
            continue
        for cell in table.rows[0].cells:
            for p in cell.paragraphs:
                if not p.paragraph_format.keep_with_next:
                    p.paragraph_format.keep_with_next = True
                    fixed += 1
    return fixed


def fix_operation_line_keep_lines(doc) -> int:
    fixed = 0
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text or len(text) > OP_LINE_MAX_LEN:
            continue
        if text.startswith("镜 ") and "操作顺序" in text:
            if not p.paragraph_format.keep_together:
                p.paragraph_format.keep_together = True
                fixed += 1
    return fixed


def fix(path: str) -> tuple[int, int, int, int]:
    doc = Document(path)
    rows = fix_tables(doc)
    orphans = fix_heading_orphans(doc)
    header_rows = fix_header_keepnext(doc)
    op_lines = fix_operation_line_keep_lines(doc)
    doc.save(path)
    return rows, orphans, header_rows, op_lines


def selftest() -> int:
    import tempfile

    from docx import Document as _Doc
    from docx.oxml.ns import qn

    ok = True

    def expect(name, cond, extra=""):
        nonlocal ok
        print(("PASS" if cond else "FAIL"), name, extra)
        if not cond:
            ok = False

    doc = _Doc()
    p_title = doc.add_paragraph()
    p_title.add_run("LOOK 01 · 标题行").bold = True          # 整段加粗短标题 → keepNext
    p_label = doc.add_paragraph("中文标签：")                 # 全角冒号短标签 → keepNext
    p_op = doc.add_paragraph()
    p_op.add_run("镜 1 操作顺序").bold = True                 # 仅段首加粗 → keepLines
    p_op.add_run("：①上传素材；②选时长；③生成。")
    doc.add_paragraph("普通正文段落。")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "表头A"
    table.cell(0, 1).text = "表头B"
    table.cell(1, 0).text = "数据1"
    table.cell(1, 1).text = "数据2"

    tmp = tempfile.mkdtemp(prefix="fixpage_") + "/st.docx"
    doc.save(tmp)
    rows, orphans, header_rows, op_lines = fix(tmp)
    expect("表格行加 cantSplit", rows >= 1, f"rows={rows}")
    expect("加粗标题/标签加 keepNext", orphans >= 2, f"orphans={orphans}")
    expect("表头行加 keepNext", header_rows >= 1, f"header={header_rows}")
    expect("操作顺序段加 keepLines", op_lines == 1, f"op_lines={op_lines}")

    d2 = _Doc(tmp)
    tr = d2.tables[0].rows[0]._tr
    expect("cantSplit 落盘", tr.trPr.find(qn("w:cantSplit")) is not None)
    expect("tblHeader 落盘", tr.trPr.find(qn("w:tblHeader")) is not None)
    keepnext = [p for p in d2.paragraphs if p.paragraph_format.keep_with_next]
    keeptogether = [p for p in d2.paragraphs if p.paragraph_format.keep_together]
    expect("keepNext 段落数 ≥2（标题/标签）", len(keepnext) >= 2, str(len(keepnext)))
    cell_p = d2.tables[0].rows[0].cells[0].paragraphs[0]
    expect("表头单元格 keepNext", bool(cell_p.paragraph_format.keep_with_next))
    expect("keepLines 段落为操作顺序段",
           any(p.text.startswith("镜 1 操作顺序") for p in keeptogether),
           str([p.text[:12] for p in keeptogether]))

    print("SELFTEST", "PASS" if ok else "FAIL")
    return 0 if ok else 1


if __name__ == "__main__":
    if "--selftest" in sys.argv:
        sys.exit(selftest())
    if len(sys.argv) != 2:
        print(__doc__)
        sys.exit(2)
    rows, orphans, header_rows, op_lines = fix(sys.argv[1])
    print(f"fixed: table_rows={rows}, heading_orphans_keepnext={orphans}, header_keepnext={header_rows}, op_line_keep_lines={op_lines} -> {sys.argv[1]}")
